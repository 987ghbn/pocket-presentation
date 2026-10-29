"""Компьютерная версия: python make_pptx.py doklad.txt [макс_слайдов]
Установка: pip install python-pptx"""
import re, sys
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

def short(s, n=110):
    s = s.strip()
    if len(s) <= n: return s
    c = s[:n-3]; i = c.rfind(" ")
    return (c[:i] if i > 40 else c) + "…"

def sents(t): return [s.strip() for s in re.split(r"(?<=[.!?…])\s+", t) if len(s.strip()) > 2]

def is_head(l, nxt):
    if re.match(r"^#{1,3}\s", l): return True
    if len(l) > 80: return False
    if re.match(r"^(\d+[.)]|[IVX]+\.|Глава|Раздел|Часть)\s", l, re.I): return True
    return len(l) <= 60 and not re.search(r"[.!?:;,]$", l) and bool(nxt) and len(nxt) > len(l)

def clean(l): return re.sub(r"^(#{1,3}\s*|(\d+[.)]|[IVX]+\.)\s+)", "", l)

def parse(raw, mx):
    lines = [l.strip() for l in raw.splitlines() if l.strip()]
    title = clean(lines[0]) if len(clean(lines[0])) <= 90 else "Презентация"
    if clean(lines[0]) == title: lines = lines[1:]
    secs, cur = [], None
    for i, l in enumerate(lines):
        nxt = lines[i+1] if i+1 < len(lines) else None
        if is_head(l, nxt): cur = {"t": clean(l), "b": []}; secs.append(cur)
        else:
            if cur is None: cur = {"t": "", "b": []}; secs.append(cur)
            cur["b"].append(l)
    if len([s for s in secs if s["t"]]) < 2:
        secs = [{"t": " ".join(p.split()[:6]), "b": [p]} for p in lines]
    limit = 99
    while True:
        out = []
        for s in secs:
            b = [short(x) for x in sents(" ".join(s["b"]))[:limit]]
            for i in range(0, max(len(b), 1), 5):
                out.append(((s["t"] or "Основное") + (" (прод.)" if i else ""), b[i:i+5]))
        if len(out) <= mx - 2 or limit <= 1: break
        limit = 3 if limit > 3 else limit - 1
    return title, out[:mx-2]

def main():
    src = sys.argv[1]; mx = int(sys.argv[2]) if len(sys.argv) > 2 else 12
    title, slides = parse(open(src, encoding="utf-8").read(), mx)
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text = title
    s.placeholders[1].text = ""
    for t, bullets in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = t
        tf = s.placeholders[1].text_frame; tf.clear()
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b; p.font.size = Pt(24)
    s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = "Спасибо за внимание!"
    out = re.sub(r"\.\w+$", "", src) + ".pptx"; prs.save(out); print("Готово:", out)

main()
